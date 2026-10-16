#!/usr/bin/env python3
"""
Builds app/src/main/assets/assigned-exchanges.bin from NANPA's public code reports.

NANPA lists every exchange (the middle three digits) of every US area code as assigned,
vacant or unavailable. A call from an exchange that isn't assigned can't be real: someone
made the number up. The app uses this file to spot those calls, entirely offline.

    pip install openpyxl
    python3 scripts/build-exchanges.py

The file is plain bits, so the phone never parses anything:
  "NEWTEXCH" (8 bytes) · format 1 (4 bytes) · data date YYYYMMDD (4 bytes)
  · area codes covered (1000 bits) · assigned exchanges (1,000,000 bits, index NPA*1000+NXX)
  · CRC-32 of everything before it (4 bytes). Integers are big-endian.
"""
import io
import re
import struct
import sys
import urllib.request
import zipfile
import zlib
from datetime import datetime
from pathlib import Path

import openpyxl

REGIONS = ["CENTRAL", "EASTERN", "WESTERN"]
URL = "https://reports.nanpa.com/public/CoCodeAssignment_Region_{}.zip"
# Assigned, protected during a split, or reserved by a carrier: all can carry real calls.
REAL = {"AS", "PR", "RV"}
OUT = Path(__file__).resolve().parent.parent / "app/src/main/assets/assigned-exchanges.bin"


def fetch(region: str) -> bytes:
    request = urllib.request.Request(URL.format(region), headers={"User-Agent": "Mozilla/5.0 (newt-phone build)"})
    with urllib.request.urlopen(request, timeout=120) as response:
        return response.read()


def main() -> None:
    npas = bytearray(1000 // 8)
    assigned = bytearray(1_000_000 // 8)
    dates, rows, real = set(), 0, 0
    for region in REGIONS:
        archive = zipfile.ZipFile(io.BytesIO(fetch(region)))
        name = next(n for n in archive.namelist() if n.upper().endswith(".XLSX"))
        book = openpyxl.load_workbook(io.BytesIO(archive.read(name)), read_only=True)
        for sheet in book.worksheets:
            if sheet.title == "UseCodes":
                continue
            header = next(sheet.iter_rows(max_row=1, values_only=True))
            found = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(header[5] or ""))
            if found:
                dates.add(int(f"{found[3]}{int(found[1]):02d}{int(found[2]):02d}"))
            for row in sheet.iter_rows(min_row=4, values_only=True):
                code = str(row[0] or "").replace("-", "")
                if not re.fullmatch(r"[2-9]\d\d[2-9]\d\d", code):
                    continue
                rows += 1
                npa, index = int(code[:3]), int(code)
                npas[npa // 8] |= 1 << (npa % 8)
                if row[1] in REAL:
                    real += 1
                    assigned[index // 8] |= 1 << (index % 8)

    area_codes = sum(bin(b).count("1") for b in npas)
    # Refuse to write a file that looks wrong; an empty one would mark every caller fake.
    if area_codes < 300 or real < 150_000 or rows < 250_000:
        sys.exit(f"Refusing to write: {area_codes} area codes, {real} assigned of {rows} rows")

    date = max(dates) if dates else int(datetime.now().strftime("%Y%m%d"))
    body = b"NEWTEXCH" + struct.pack(">II", 1, date) + bytes(npas) + bytes(assigned)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF))
    print(f"{OUT.name}: data from {date}, {area_codes} area codes, {real} of {rows} exchanges assigned")


if __name__ == "__main__":
    main()
